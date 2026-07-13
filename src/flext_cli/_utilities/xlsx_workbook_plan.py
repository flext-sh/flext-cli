"""Prepare an exact workbook surface from one validated XLSX plan."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.worksheet import Worksheet

from flext_cli import p, r
from flext_cli._constants.xlsx import FlextCliConstantsXlsx
from flext_cli._models.xlsx_workbook import FlextCliModelsXlsxWorkbook

from .xlsx_style_codec import FlextCliUtilitiesXlsxStyleCodec
from .xlsx_workbook_io import FlextCliUtilitiesXlsxWorkbookIo


class FlextCliUtilitiesXlsxWorkbookPlan(
    FlextCliUtilitiesXlsxStyleCodec, FlextCliUtilitiesXlsxWorkbookIo
):
    """Load visual resources and recreate the exact planned sheet graph."""

    # NOTE (multi-agent, mro-j2yt.1): template sheets and names are discarded;
    # only visual resources survive, so stale document content cannot leak.
    @staticmethod
    def _validate_plan(
        plan: FlextCliModelsXlsxWorkbook.XlsxWorkbookPlan,
    ) -> p.Result[bool]:
        sheet_names: frozenset[str] = frozenset()
        style_names: frozenset[str] = frozenset()
        defined_names: frozenset[str] = frozenset()
        table_names: frozenset[str] = frozenset()
        for sheet in plan.sheets:
            if sheet.name in sheet_names:
                return r[bool].fail(
                    f"{FlextCliConstantsXlsx.XlsxError.DUPLICATE_SHEET}: {sheet.name}"
                )
            sheet_names = sheet_names.union((sheet.name,))
            for table in sheet.tables:
                if table.name in table_names:
                    return r[bool].fail(
                        f"{FlextCliConstantsXlsx.XlsxError.DUPLICATE_TABLE}: "
                        f"{table.name}"
                    )
                table_names = table_names.union((table.name,))
        for style in plan.named_styles:
            if style.name in style_names:
                return r[bool].fail(f"Duplicate named style in plan: {style.name}")
            style_names = style_names.union((style.name,))
        for item in plan.defined_names:
            if item.name in defined_names:
                return r[bool].fail(
                    f"{FlextCliConstantsXlsx.XlsxError.DUPLICATE_DEFINED_NAME}: "
                    f"{item.name}"
                )
            defined_names = defined_names.union((item.name,))
        return r[bool].ok(True)

    @classmethod
    def _workbook_for_request(
        cls, request: FlextCliModelsXlsxWorkbook.XlsxRenderRequest
    ) -> p.Result[Workbook]:
        validation = cls._validate_plan(request.plan)
        if validation.failure:
            return r[Workbook].fail(
                validation.error or str(FlextCliConstantsXlsx.XlsxError.PLAN_INVALID)
            )
        if request.template is None:
            workbook = cls._new_workbook()
        else:
            loaded = cls._load_workbook(request.template)
            if loaded.failure:
                return r[Workbook].fail(
                    loaded.error
                    or str(FlextCliConstantsXlsx.XlsxError.WORKBOOK_LOAD_FAILED)
                )
            workbook = loaded.value
        try:
            for worksheet in tuple(workbook.worksheets):
                workbook.remove(worksheet)
            for chartsheet in tuple(workbook.chartsheets):
                workbook.remove(chartsheet)
            workbook.defined_names.clear()
            for spec in request.plan.named_styles:
                if spec.name not in workbook.named_styles:
                    workbook.add_named_style(cls._named_style(spec))
            for sheet in request.plan.sheets:
                created = workbook.create_sheet(sheet.name)
                if not isinstance(created, Worksheet):
                    return r[Workbook].fail(f"Worksheet creation failed: {sheet.name}")
            full = request.plan.full_calculation_on_load
            workbook.calculation = CalcProperties(
                calcMode=FlextCliConstantsXlsx.XLSX_DEFAULT_CALCULATION_MODE,
                fullCalcOnLoad=full,
                forceFullCalc=full,
                calcOnSave=full,
            )
        except (KeyError, TypeError, ValueError) as exc:
            detail = str(exc).strip() or exc.__class__.__name__
            return r[Workbook].fail(
                f"{FlextCliConstantsXlsx.XlsxError.RENDER_FAILED}: {detail}"
            )
        return r[Workbook].ok(workbook)


__all__: tuple[str, ...] = ("FlextCliUtilitiesXlsxWorkbookPlan",)
